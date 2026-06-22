"""Excel export — one workbook per scenario (Budget, Liquidität, Personal,
Einnahmen, Ausgaben), and a scenario-comparison workbook focused on differences.

German throughout, with a date/time stamp and charts. Uses openpyxl.
"""
from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlmodel import Session, select

from ..db import get_settings
from ..engine.liquidity import liquidity_view, liquidity_view_for_scenario
from ..engine.payroll_at import employee_year_cost, ruleset_for_year
from ..engine.pnl import pnl_view
from ..engine.scenarios import effective_cost_ids, effective_revenue_ids, get_scenario, list_scenarios
from ..models import (
    CostCategory,
    CostPlanMonth,
    Employee,
    RevenuePlanMonth,
    RevenueStream,
    SalaryMonth,
)
from ..ui.formatting import MONTHS_DE, YEARS

_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="2563EB")
_TITLE = Font(bold=True, size=14)
_SUB = Font(italic=True, color="666666", size=9)
_BOLD = Font(bold=True)
_EUR = '#,##0" €"'
_THIN = Side(style="thin", color="DDDDDD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _title(ws, title: str, subtitle: str) -> None:
    ws["A1"] = title
    ws["A1"].font = _TITLE
    ws["A2"] = subtitle
    ws["A2"].font = _SUB


def _header_row(ws, row: int, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = _HEAD
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = _BORDER


def _autosize(ws, widths: dict[int, int]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


def _month_headers(months: list[int]) -> list[str]:
    total = "Jahr" if len(months) == 12 else "Summe"
    return ["Position", *[MONTHS_DE[m - 1] for m in months], total]


def _year_months(start: tuple[int, int], end: tuple[int, int]) -> dict[int, list[int]]:
    """Ordered {year: [months]} for the inclusive (year, month) range."""
    out: dict[int, list[int]] = {}
    y, m = start
    while (y, m) <= end:
        out.setdefault(y, []).append(m)
        m += 1
        if m > 12:
            m, y = 1, y + 1
    return out


# --- data helpers ----------------------------------------------------------------

def _revenue_matrix(session: Session, scenario, year: int):
    ids = effective_revenue_ids(session, scenario)
    streams = [st for st in session.exec(select(RevenueStream).order_by(
        RevenueStream.sort_order, RevenueStream.id)).all() if st.id in ids]
    out = []
    for st in streams:
        months = {p.month: p.amount for p in session.exec(select(RevenuePlanMonth).where(
            RevenuePlanMonth.stream_id == st.id, RevenuePlanMonth.year == year)).all()}
        vals = [round(months.get(m, 0.0)) for m in range(1, 13)]
        out.append((st.name, vals))
    return out


def _cost_matrix(session: Session, scenario, year: int):
    ids = effective_cost_ids(session, scenario)
    cats = [c for c in session.exec(select(CostCategory).order_by(
        CostCategory.sort_order, CostCategory.id)).all() if c.id in ids]
    out = []
    for c in cats:
        months = {p.month: p.amount for p in session.exec(select(CostPlanMonth).where(
            CostPlanMonth.category_id == c.id, CostPlanMonth.year == year)).all()}
        vals = [round(months.get(m, 0.0)) for m in range(1, 13)]
        out.append((c.name, vals))
    return out


def _personnel_rows(session: Session, year: int):
    settings = get_settings(session)
    rs = ruleset_for_year(year)
    rows = []
    for e in session.exec(select(Employee).order_by(Employee.sort_order, Employee.id)).all():
        regular = [0.0] * 12
        special = [0.0] * 12
        for sm in session.exec(select(SalaryMonth).where(
                SalaryMonth.employee_id == e.id, SalaryMonth.year == year)).all():
            regular[sm.month - 1] += sm.gross
            special[sm.month - 1] += sm.special
        if not any(regular) and not any(special):
            continue
        if e.is_contractor:
            gross = [regular[m] + special[m] for m in range(12)]
            rows.append((e.name + " (extern)", gross, gross, gross))
        else:
            mc = employee_year_cost(regular, special, burden_pct=settings.employer_burden_pct, rs=rs)
            rows.append((e.name, [round(x) for x in mc.gross], [round(x) for x in mc.net],
                         [round(x) for x in mc.all_in]))
    return rows


# --- sheet builders --------------------------------------------------------------

def _sheet_budget(ws, session: Session, scenario_id: int, stamp: str, sc_name: str,
                  year_months: dict[int, list[int]]) -> None:
    _title(ws, "GuV / Budget (Deckungsbeitragsrechnung)",
           f"Szenario: {sc_name} · Erstellt: {stamp}")
    row = 4
    for year, months in year_months.items():
        last = 2 + len(months)
        ws.cell(row=row, column=1, value=f"Jahr {year}").font = _BOLD
        row += 1
        _header_row(ws, row, _month_headers(months))
        row += 1
        for pr in pnl_view(session, year, scenario_id):
            ws.cell(row=row, column=1, value=pr.label.strip())
            for i, m in enumerate(months):
                c = ws.cell(row=row, column=2 + i, value=round(pr.values[m - 1]))
                c.number_format = _EUR
            jc = ws.cell(row=row, column=last, value=round(sum(pr.values[m - 1] for m in months)))
            jc.number_format = _EUR
            if pr.kind == "subtotal":
                for col in range(1, last + 1):
                    ws.cell(row=row, column=col).font = _BOLD
            row += 1
        row += 2
    _autosize(ws, {1: 32, **{c: 12 for c in range(2, 14)}, 14: 14})


def _sheet_liquidity(ws, session: Session, scenario_id: int, stamp: str, sc_name: str,
                     start_date=None, end_date=None) -> None:
    _title(ws, "Liquidität", f"Szenario: {sc_name} · Erstellt: {stamp}")
    headers = ["Datum", "Einzahlungen", "Auszahlungen", "davon Subunternehmer",
               "Saldo", "Bank Status", "ohne EU-Förderung"]
    _header_row(ws, 4, headers)
    rows = liquidity_view_for_scenario(session, scenario_id)
    if start_date or end_date:
        rows = [br for br in rows
                if (start_date is None or br.bucket >= start_date)
                and (end_date is None or br.bucket <= end_date)]
    r = 5
    for br in rows:
        ws.cell(row=r, column=1, value=br.label)
        for col, val in enumerate([br.inflow, br.outflow, br.subcontractor, br.net,
                                    br.balance, br.balance_no_eu], start=2):
            c = ws.cell(row=r, column=col, value=round(val))
            c.number_format = _EUR
        r += 1
    _autosize(ws, {1: 12, 2: 14, 3: 14, 4: 18, 5: 13, 6: 14, 7: 16})
    # Line chart of the running balance.
    if rows:
        chart = LineChart()
        chart.title = "Bank Status"
        chart.height = 9
        chart.width = 24
        data = Reference(ws, min_col=6, min_row=4, max_col=7, max_row=4 + len(rows))
        cats = Reference(ws, min_col=1, min_row=5, max_row=4 + len(rows))
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        ws.add_chart(chart, f"I4")


def _sheet_personal(ws, session: Session, stamp: str, sc_name: str,
                    year_months: dict[int, list[int]]) -> None:
    _title(ws, "Personal", f"Erstellt: {stamp}")
    row = 4
    for year, months in year_months.items():
        last = 2 + len(months)
        ws.cell(row=row, column=1, value=f"Jahr {year}").font = _BOLD
        row += 1
        _header_row(ws, row, _month_headers(months))
        row += 1
        emps = _personnel_rows(session, year)
        sums = {0: [0.0] * 12, 1: [0.0] * 12, 2: [0.0] * 12}  # gross / net / all-in
        for name, gross, net, allin in emps:
            ws.cell(row=row, column=1, value=name).font = _BOLD
            row += 1
            for idx, (label, arr) in enumerate((("  Brutto", gross), ("  Netto", net),
                                                ("  Gesamtkosten (Brutto+LNK)", allin))):
                ws.cell(row=row, column=1, value=label)
                for i, m in enumerate(months):
                    c = ws.cell(row=row, column=2 + i, value=round(arr[m - 1]))
                    c.number_format = _EUR
                    sums[idx][m - 1] += arr[m - 1]
                jc = ws.cell(row=row, column=last, value=round(sum(arr[m - 1] for m in months)))
                jc.number_format = _EUR
                row += 1
        # Summary across all employees.
        if emps:
            ws.cell(row=row, column=1, value="Summe alle Mitarbeiter").font = _BOLD
            row += 1
            for idx, label in ((0, "  Brutto gesamt"), (1, "  Netto gesamt"),
                               (2, "  Gesamtkosten gesamt (Brutto+LNK)")):
                lc = ws.cell(row=row, column=1, value=label)
                lc.font = _BOLD
                for i, m in enumerate(months):
                    c = ws.cell(row=row, column=2 + i, value=round(sums[idx][m - 1]))
                    c.number_format = _EUR
                    c.font = _BOLD
                jc = ws.cell(row=row, column=last,
                             value=round(sum(sums[idx][m - 1] for m in months)))
                jc.number_format = _EUR
                jc.font = _BOLD
                row += 1
        row += 2
    _autosize(ws, {1: 32, **{c: 12 for c in range(2, 14)}, 14: 14})


def _sheet_matrix(ws, title: str, stamp: str, sc_name: str, session: Session,
                  data_fn, scenario, year_months: dict[int, list[int]]) -> None:
    _title(ws, title, f"Szenario: {sc_name} · Erstellt: {stamp}")
    row = 4
    for year, months in year_months.items():
        last = 2 + len(months)
        ws.cell(row=row, column=1, value=f"Jahr {year}").font = _BOLD
        row += 1
        _header_row(ws, row, _month_headers(months))
        row += 1
        total = [0] * 12
        for name, vals in data_fn(session, scenario, year):
            ws.cell(row=row, column=1, value=name)
            for i, m in enumerate(months):
                c = ws.cell(row=row, column=2 + i, value=vals[m - 1])
                c.number_format = _EUR
                total[m - 1] += vals[m - 1]
            jc = ws.cell(row=row, column=last, value=sum(vals[m - 1] for m in months))
            jc.number_format = _EUR
            row += 1
        ws.cell(row=row, column=1, value="Summe").font = _BOLD
        for i, m in enumerate(months):
            c = ws.cell(row=row, column=2 + i, value=total[m - 1])
            c.number_format = _EUR
            c.font = _BOLD
        c = ws.cell(row=row, column=last, value=sum(total[m - 1] for m in months))
        c.number_format = _EUR
        c.font = _BOLD
        row += 3
    _autosize(ws, {1: 28, **{c: 12 for c in range(2, 14)}, 14: 14})


# --- public API ------------------------------------------------------------------

def export_single_scenario(session: Session, scenario_id: int = 1,
                           start: tuple[int, int] | None = None,
                           end: tuple[int, int] | None = None) -> bytes:
    """Export one scenario. `start`/`end` are (year, month) bounds (inclusive);
    default is the full horizon (all configured years)."""
    from calendar import monthrange
    from datetime import date

    settings = get_settings(session)
    scenario = get_scenario(session, scenario_id)
    stamp = datetime.now().strftime("%d.%m.%Y %H:%M")
    company = settings.company_name

    start = start or (YEARS[0], 1)
    end = end or (YEARS[-1], 12)
    if end < start:
        start, end = end, start
    ym = _year_months(start, end)
    start_date = date(start[0], start[1], 1)
    end_date = date(end[0], end[1], monthrange(end[0], end[1])[1])

    wb = Workbook()
    wb.properties.title = f"Budget & Liquidität — {company} — {scenario.name}"

    ws = wb.active
    ws.title = "Budget"
    _sheet_budget(ws, session, scenario_id, stamp, scenario.name, ym)
    _sheet_liquidity(wb.create_sheet("Liquidität"), session, scenario_id, stamp,
                     scenario.name, start_date, end_date)
    _sheet_personal(wb.create_sheet("Personal"), session, stamp, scenario.name, ym)
    _sheet_matrix(wb.create_sheet("Einnahmen"), "Einnahmen", stamp, scenario.name,
                  session, _revenue_matrix, scenario, ym)
    _sheet_matrix(wb.create_sheet("Ausgaben"), "Ausgaben", stamp, scenario.name,
                  session, _cost_matrix, scenario, ym)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_scenario_comparison(session: Session, year: int) -> bytes:
    """A workbook focused on the differences between scenarios."""
    from .snapshots import compute_kpis
    settings = get_settings(session)
    stamp = datetime.now().strftime("%d.%m.%Y %H:%M")
    scs = list_scenarios(session)
    wb = Workbook()

    ws = wb.active
    ws.title = "Szenarienvergleich"
    _title(ws, f"Szenarienvergleich {year}",
           f"{settings.company_name} · Erstellt: {stamp}")
    kpis = {sc.id: compute_kpis(session, year, sc.id) for sc in scs}
    keys = list(next(iter(kpis.values())).keys()) if kpis else []
    headers = ["Kennzahl", *[sc.name for sc in scs]]
    if len(scs) >= 2:
        headers.append(f"Δ ({scs[1].name} − {scs[0].name})")
    _header_row(ws, 4, headers)
    r = 5
    for k in keys:
        ws.cell(row=r, column=1, value=k)
        for col, sc in enumerate(scs, start=2):
            c = ws.cell(row=r, column=col, value=kpis[sc.id].get(k, 0))
            c.number_format = _EUR
        if len(scs) >= 2:
            diff = kpis[scs[1].id].get(k, 0) - kpis[scs[0].id].get(k, 0)
            dc = ws.cell(row=r, column=2 + len(scs), value=diff)
            dc.number_format = _EUR
            dc.font = Font(bold=True, color=("15803D" if diff > 0 else "B91C1C"))
        r += 1
    _autosize(ws, {1: 26, **{c: 16 for c in range(2, 2 + len(scs) + 1)}})

    # Bar chart: KPI value per scenario.
    chart = BarChart()
    chart.type = "col"
    chart.title = f"Kennzahlen je Szenario ({year})"
    chart.height = 10
    chart.width = 26
    data = Reference(ws, min_col=2, min_row=4, max_col=1 + len(scs), max_row=4 + len(keys))
    cats = Reference(ws, min_col=1, min_row=5, max_row=4 + len(keys))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    ws.add_chart(chart, f"A{r + 2}")

    # Liquidity overlay sheet: Bank Status per scenario over time.
    ws2 = wb.create_sheet("Liquidität-Vergleich")
    _title(ws2, f"Liquidität je Szenario {year}", f"Erstellt: {stamp}")
    series = []
    for sc in scs:
        liq = (liquidity_view(session) if sc.base_id is None
               else liquidity_view_for_scenario(session, sc.id))
        series.append((sc.name, liq))
    if series:
        labels = [b.label for b in series[0][1]]
        ws2.cell(row=4, column=1, value="Datum").font = _HEAD
        ws2.cell(row=4, column=1).fill = _HEAD_FILL
        for i, (nm, _) in enumerate(series, start=2):
            cell = ws2.cell(row=4, column=i, value=nm)
            cell.font = _HEAD
            cell.fill = _HEAD_FILL
        for rr, lab in enumerate(labels, start=5):
            ws2.cell(row=rr, column=1, value=lab)
            for i, (_, liq) in enumerate(series, start=2):
                bal = liq[rr - 5].balance if rr - 5 < len(liq) else 0
                c = ws2.cell(row=rr, column=i, value=round(bal))
                c.number_format = _EUR
        chart2 = LineChart()
        chart2.title = "Bank Status je Szenario"
        chart2.height = 10
        chart2.width = 28
        data2 = Reference(ws2, min_col=2, min_row=4, max_col=1 + len(series),
                          max_row=4 + len(labels))
        cats2 = Reference(ws2, min_col=1, min_row=5, max_row=4 + len(labels))
        chart2.add_data(data2, titles_from_data=True)
        chart2.set_categories(cats2)
        ws2.add_chart(chart2, f"{get_column_letter(len(series) + 3)}4")
        _autosize(ws2, {1: 12, **{c: 16 for c in range(2, 2 + len(series))}})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
